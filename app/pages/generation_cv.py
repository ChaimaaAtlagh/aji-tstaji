"""
generation_cv.py
Page « Génération CV » : crée un CV (Excel, openpyxl) et un message de
candidature (texte) à partir du profil (SQLite). Aucune donnée inventée.
"""

import tkinter as tk
from pathlib import Path
from tkinter import messagebox, ttk

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from app import theme
from app.widgets import BoutonModerne
from src.database.base_donnees import initialiser_base, obtenir_profil

RACINE = Path(__file__).resolve().parents[2]
DOSSIER_SORTIE = RACINE / "data" / "reports" / "cv_genere"

VIDE = "-"
BLEU_HEX = "2563EB"
MARINE_HEX = "0F1B4C"


def texte(valeur):
    if valeur is None or str(valeur).strip() in ("", "nan"):
        return VIDE
    return str(valeur)


class GenerationCV(tk.Frame):
    def __init__(self, parent, app):
        super().__init__(parent, bg=theme.FOND)
        self.app = app
        initialiser_base()

        self._creer_defilement()
        self._creer_titre()
        self.zone_apercu = tk.Frame(self.contenu, bg=theme.FOND)
        self.zone_apercu.pack(fill="both", expand=True, padx=50)
        self._creer_actions()

    def _creer_defilement(self):
        self.canvas = tk.Canvas(self, bg=theme.FOND, highlightthickness=0, yscrollincrement=30)
        barre = ttk.Scrollbar(self, orient="vertical", command=self.canvas.yview)
        self.canvas.configure(yscrollcommand=barre.set)
        barre.pack(side="right", fill="y")
        self.canvas.pack(side="left", fill="both", expand=True)

        self.contenu = tk.Frame(self.canvas, bg=theme.FOND)
        fenetre = self.canvas.create_window((0, 0), window=self.contenu, anchor="nw")
        self.contenu.bind("<Configure>",
            lambda e: self.canvas.configure(scrollregion=self.canvas.bbox("all")))
        self.canvas.bind("<Configure>", lambda e: self.canvas.itemconfig(fenetre, width=e.width))
        self.bind_all("<MouseWheel>", self._molette, add="+")

    def _molette(self, evenement):
        if self.winfo_ismapped():
            self.canvas.yview_scroll(int(-evenement.delta / 120), "units")

    def _creer_titre(self):
        tk.Label(self.contenu, text="Génération CV", font=theme.TITRE, fg=theme.MARINE,
                 bg=theme.FOND).pack(anchor="w", padx=50, pady=(25, 0))
        tk.Label(self.contenu, text="Génère ton CV (Excel) et un message de candidature "
                            "à partir des informations de ton profil.",
                 font=theme.NORMAL, fg=theme.TEXTE_CLAIR,
                 bg=theme.FOND).pack(anchor="w", padx=50, pady=(0, 20))

    def actualiser(self):
        for enfant in self.zone_apercu.winfo_children():
            enfant.destroy()

        profil = obtenir_profil()
        if not profil or not profil.get("nom") or not profil.get("email"):
            carte = tk.Frame(self.zone_apercu, bg=theme.BLANC,
                             highlightbackground=theme.BORDURE, highlightthickness=1)
            carte.pack(fill="x", pady=10)
            contenu = tk.Frame(carte, bg=theme.BLANC)
            contenu.pack(padx=24, pady=24)
            tk.Label(contenu, text="Ton profil n'est pas encore complet",
                     font=(theme.POLICE, 13, "bold"), fg=theme.MARINE, bg=theme.BLANC
                     ).pack(anchor="w")
            tk.Label(contenu, text="Renseigne au moins ton nom et ton email dans "
                                   "« Mon profil » pour générer un CV.",
                     font=theme.NORMAL, fg=theme.TEXTE_CLAIR, bg=theme.BLANC
                     ).pack(anchor="w", pady=(6, 14))
            BoutonModerne(contenu, texte="Aller à Mon profil", icone="→", style="principal",
                         commande=lambda: self.app.afficher_page("Mon profil"),
                         fond_parent=theme.BLANC, hauteur=36).pack(anchor="w")
            return

        self._carte_apercu(profil)

    def _carte_apercu(self, profil):
        carte = tk.Frame(self.zone_apercu, bg=theme.BLANC,
                         highlightbackground=theme.BORDURE, highlightthickness=1)
        carte.pack(fill="x", pady=10)
        contenu = tk.Frame(carte, bg=theme.BLANC)
        contenu.pack(fill="x", padx=22, pady=18)

        nom_complet = f"{texte(profil.get('prenom'))} {texte(profil.get('nom'))}".strip()
        tk.Label(contenu, text=nom_complet, font=(theme.POLICE, 14, "bold"), fg=theme.MARINE,
                 bg=theme.BLANC).pack(anchor="w")
        tk.Label(contenu,
                 text=f"{texte(profil.get('diplome'))} · {texte(profil.get('etablissement'))}",
                 font=theme.NORMAL, fg=theme.TEXTE_CLAIR, bg=theme.BLANC
                 ).pack(anchor="w", pady=(2, 10))

        infos = [("Email", profil.get("email")), ("Téléphone", profil.get("telephone")),
                 ("Ville", profil.get("ville")), ("Je cherche", profil.get("recherche"))]
        grille = tk.Frame(contenu, bg=theme.BLANC)
        grille.pack(anchor="w", fill="x")
        for i, (titre, valeur) in enumerate(infos):
            ligne, colonne = i // 2, i % 2
            bloc = tk.Frame(grille, bg=theme.BLANC)
            bloc.grid(row=ligne, column=colonne, sticky="w", padx=(0, 30), pady=4)
            tk.Label(bloc, text=f"{titre} :", font=theme.PETIT, fg=theme.TEXTE_CLAIR,
                     bg=theme.BLANC).pack(side="left")
            tk.Label(bloc, text=texte(valeur), font=theme.PETIT, fg=theme.TEXTE,
                     bg=theme.BLANC).pack(side="left", padx=(4, 0))

        entreprise = self.app.entreprise_choisie
        if entreprise:
            tk.Frame(contenu, bg=theme.BORDURE, height=1).pack(fill="x", pady=14)
            tk.Label(contenu, text=f"Message personnalisé pour : "
                                   f"{texte(entreprise.get('Nom_Entreprise'))}",
                     font=theme.PETIT, fg=theme.TEXTE_CLAIR, bg=theme.BLANC).pack(anchor="w")

    def _creer_actions(self):
        bas = tk.Frame(self.contenu, bg=theme.FOND)
        bas.pack(fill="x", padx=50, pady=(10, 40))
        self.message_resultat = tk.Label(bas, text="", font=theme.PETIT, fg=theme.VERT,
                                         bg=theme.FOND, wraplength=700, justify="left")
        self.message_resultat.pack(anchor="w", pady=(0, 10))

        boutons = tk.Frame(bas, bg=theme.FOND)
        boutons.pack(anchor="w")
        BoutonModerne(boutons, texte="Générer le CV (Excel)", icone="📄", style="principal",
                     commande=self.generer_cv, fond_parent=theme.FOND,
                     hauteur=38).pack(side="left", padx=(0, 10))
        BoutonModerne(boutons, texte="Générer le message", icone="✉", style="secondaire",
                     commande=self.generer_message, fond_parent=theme.FOND,
                     hauteur=38).pack(side="left")

    def generer_cv(self):
        profil = obtenir_profil()
        if not profil or not profil.get("nom"):
            messagebox.showwarning("AJI TSTAJI", "Complète d'abord ton profil.")
            return

        DOSSIER_SORTIE.mkdir(parents=True, exist_ok=True)
        classeur = Workbook()
        feuille = classeur.active
        feuille.title = "CV"
        feuille.column_dimensions["A"].width = 22
        feuille.column_dimensions["B"].width = 60

        police_titre = Font(size=18, bold=True, color=MARINE_HEX)
        police_section = Font(size=12, bold=True, color="FFFFFF")
        fond_section = PatternFill("solid", fgColor=BLEU_HEX)
        police_normale = Font(size=11)

        nom_complet = f"{texte(profil.get('prenom'))} {texte(profil.get('nom'))}".strip()
        feuille["A1"] = nom_complet
        feuille["A1"].font = police_titre
        feuille.merge_cells("A1:B1")

        contact = (f"{texte(profil.get('email'))}  ·  {texte(profil.get('telephone'))}  ·  "
                  f"{texte(profil.get('ville'))}")
        feuille["A2"] = contact
        feuille.merge_cells("A2:B2")

        ligne = 4

        def section(titre_section):
            nonlocal ligne
            feuille.cell(row=ligne, column=1, value=titre_section).font = police_section
            feuille.cell(row=ligne, column=1).fill = fond_section
            feuille.cell(row=ligne, column=2).fill = fond_section
            feuille.merge_cells(start_row=ligne, start_column=1, end_row=ligne, end_column=2)
            ligne += 1

        def champ(titre_champ, valeur):
            nonlocal ligne
            feuille.cell(row=ligne, column=1, value=titre_champ).font = Font(bold=True, size=11)
            cellule = feuille.cell(row=ligne, column=2, value=texte(valeur))
            cellule.font = police_normale
            cellule.alignment = Alignment(wrap_text=True, vertical="top")
            ligne += 1

        section("Formation")
        champ("Diplôme préparé", profil.get("diplome"))
        champ("Secteur étudié", profil.get("secteur"))
        champ("Établissement", profil.get("etablissement"))

        ligne += 1
        section("Recherche")
        champ("Je cherche", profil.get("recherche"))

        ligne += 1
        section("Compétences")
        feuille.cell(row=ligne, column=1, value="Compétences").font = Font(bold=True, size=11)
        cellule = feuille.cell(row=ligne, column=2, value=texte(profil.get("competences")))
        cellule.alignment = Alignment(wrap_text=True, vertical="top")
        feuille.row_dimensions[ligne].height = 60
        ligne += 1

        ligne += 1
        section("Liens")
        champ("LinkedIn", profil.get("linkedin"))

        nom_fichier = f"CV_{texte(profil.get('nom')).replace(' ', '_')}.xlsx"
        chemin = DOSSIER_SORTIE / nom_fichier
        classeur.save(chemin)

        self.message_resultat.config(text=f"CV généré : {chemin}", fg=theme.VERT)

    def generer_message(self):
        profil = obtenir_profil()
        entreprise = self.app.entreprise_choisie

        if not profil or not profil.get("nom"):
            messagebox.showwarning("AJI TSTAJI", "Complète d'abord ton profil.")
            return
        if not entreprise:
            messagebox.showwarning("AJI TSTAJI",
                "Choisis d'abord une entreprise depuis « Entreprises » ou "
                "« Opportunités » (bouton Postuler) pour un message personnalisé.")
            return

        DOSSIER_SORTIE.mkdir(parents=True, exist_ok=True)
        nom_complet = f"{texte(profil.get('prenom'))} {texte(profil.get('nom'))}".strip()
        nom_entreprise = texte(entreprise.get("Nom_Entreprise"))

        message = (
            f"Objet : Candidature - {texte(profil.get('recherche'))}\n\n"
            f"Bonjour,\n\n"
            f"Je me permets de vous contacter dans le cadre de ma recherche "
            f"de {texte(profil.get('recherche')).lower()}. "
            f"Actuellement en {texte(profil.get('diplome'))} "
            f"à {texte(profil.get('etablissement'))}, je suis particulièrement "
            f"intéressé(e) par les activités de {nom_entreprise} "
            f"dans le secteur {texte(entreprise.get('Domaine')).lower()}.\n\n"
            f"Compétences : {texte(profil.get('competences'))}\n\n"
            f"Je reste à votre disposition pour un entretien et vous remercie "
            f"de l'attention portée à ma candidature.\n\n"
            f"Cordialement,\n{nom_complet}\n"
            f"{texte(profil.get('email'))}  -  {texte(profil.get('telephone'))}\n"
        )

        nom_fichier = f"Message_{nom_entreprise.replace(' ', '_')}.txt"
        chemin = DOSSIER_SORTIE / nom_fichier
        chemin.write_text(message, encoding="utf-8")

        self.message_resultat.config(text=f"Message généré : {chemin}", fg=theme.VERT)