"""Page chatbot IA de AJI TSTAJI — connexion à Groq (design panneau demi-page)."""

import os
import threading
import customtkinter as ctk
from app.markdown import afficher_markdown
from PIL import Image, ImageDraw
from tkinter import filedialog
from dotenv import load_dotenv
from pypdf import PdfReader
from docx import Document

from pathlib import Path

RACINE_PROJET = Path(__file__).resolve().parents[1]
load_dotenv(RACINE_PROJET / ".env")

try:
    from openai import OpenAI
except ImportError:
    OpenAI = None


NOM_CHATBOT = "assistantIA-AJI TSTAJI"
MODELE_PAR_DEFAUT = "openai/gpt-oss-20b"
PLACEHOLDER = "Écris ton message ici..."

# Couleurs du nouveau design
BLEU = "#3B4FE0"
BLEU_HOVER = "#2F40C0"
FOND = "#F4F6FB"
BLANC = "#FFFFFF"
BULLE_USER = "#DCE4FF"
TEXTE = "#1A1F36"
GRIS = "#6B7280"


def creer_icone_robot(taille=64, fond=BLEU, rayon=0.28):
    """Retourne une image PIL carrée du robot.

    taille : taille en pixels
    fond   : couleur du fond (None = transparent)
    rayon  : 0.5 = rond complet, 0.28 = carré arrondi
    """
    s = 512
    k = 2  # coordonnées pensées sur 256, multipliées par 2
    img = Image.new("RGBA", (s, s), (0, 0, 0, 0))
    d = ImageDraw.Draw(img)

    if fond:
        d.rounded_rectangle((0, 0, s - 1, s - 1), radius=int(s * rayon), fill=fond)

    blanc, marine, cyan = "#FFFFFF", "#14206B", "#5EDCFF"

    # antenne
    d.line((128 * k, 62 * k, 128 * k, 80 * k), fill=blanc, width=5 * k)
    d.ellipse((118 * k, 44 * k, 138 * k, 64 * k), fill=cyan)
    # oreilles
    d.ellipse((38 * k, 112 * k, 66 * k, 140 * k), fill=blanc)
    d.ellipse((190 * k, 112 * k, 218 * k, 140 * k), fill=blanc)
    # tête
    d.rounded_rectangle((56 * k, 80 * k, 200 * k, 172 * k), radius=38 * k, fill=blanc)
    # visière sombre
    d.rounded_rectangle((74 * k, 98 * k, 182 * k, 154 * k), radius=26 * k, fill=marine)
    # yeux
    d.ellipse((96 * k, 116 * k, 116 * k, 136 * k), fill=cyan)
    d.ellipse((140 * k, 116 * k, 160 * k, 136 * k), fill=cyan)
    # corps
    d.rounded_rectangle((90 * k, 182 * k, 166 * k, 212 * k), radius=15 * k, fill=blanc)

    return img.resize((taille, taille), Image.Resampling.LANCZOS)


class ChatbotPage(ctk.CTkFrame):
    """Panneau chatbot.

    fermer : fonction appelée par le bouton « ← Retour » (ferme le panneau).
    """

    def __init__(self, master, app=None, fermer=None, **kwargs):
        kwargs.setdefault("fg_color", FOND)
        kwargs.setdefault("corner_radius", 16)
        super().__init__(master, **kwargs)

        self.app = app
        self.fermer = fermer
        self.api_key = os.getenv("GROQ_API_KEY", "").strip()
        self.modele = os.getenv("GROQ_MODEL", MODELE_PAR_DEFAUT).strip()

        self.client = None
        self.fichier_joint = None
        self.messages = []

        # Icônes du robot (garder les références)
        self.img_avatar_grand = ctk.CTkImage(
            creer_icone_robot(96), size=(46, 46))
        self.img_avatar_petit = ctk.CTkImage(
            creer_icone_robot(64, rayon=0.5), size=(30, 30))

        if OpenAI is not None and self.api_key:
            try:
                self.client = OpenAI(
                    api_key=self.api_key,
                    base_url="https://api.groq.com/openai/v1",
                )
            except Exception:
                self.client = None

        self._creer_interface()

    # ---------------------------------------------------------
    # INTERFACE
    # ---------------------------------------------------------

    def _creer_interface(self):
        self.grid_columnconfigure(0, weight=1)
        self.grid_rowconfigure(1, weight=1)

        # ----- En-tête -----
        entete = ctk.CTkFrame(self, fg_color=BLANC, corner_radius=14)
        entete.grid(row=0, column=0, sticky="ew", padx=12, pady=(12, 6))
        entete.grid_columnconfigure(1, weight=1)

        ctk.CTkLabel(entete, image=self.img_avatar_grand, text="").grid(
            row=0, column=0, rowspan=2, padx=(12, 10), pady=10)

        ctk.CTkLabel(
            entete, text=NOM_CHATBOT, text_color=TEXTE, anchor="w",
            font=ctk.CTkFont(size=17, weight="bold"),
        ).grid(row=0, column=1, sticky="sw", pady=(10, 0))

        ctk.CTkLabel(
            entete, text="Stages, candidatures, CV et informatique",
            text_color=GRIS, anchor="w", font=ctk.CTkFont(size=11),
        ).grid(row=1, column=1, sticky="nw", pady=(0, 10))

        boutons = ctk.CTkFrame(entete, fg_color="transparent")
        boutons.grid(row=0, column=2, rowspan=2, padx=(4, 10))

        ctk.CTkButton(
            boutons, text="← Retour", width=76, height=30,
            fg_color="transparent", text_color=BLEU,
            border_width=1, border_color=BLEU, hover_color="#E8ECFF",
            font=ctk.CTkFont(size=12), command=self._retour,
        ).pack(side="left", padx=(0, 6))

        ctk.CTkButton(
            boutons, text="Nouvelle conversation", width=140, height=30,
            fg_color=BLEU, hover_color=BLEU_HOVER,
            font=ctk.CTkFont(size=12), command=self.nouvelle_conversation,
        ).pack(side="left")

        # ----- Zone des messages -----
        self.zone_messages = ctk.CTkScrollableFrame(
            self, fg_color=FOND, corner_radius=10)
        self.zone_messages.grid(
            row=1, column=0, sticky="nsew", padx=8, pady=(0, 4))
        self.zone_messages.grid_columnconfigure(0, weight=1)

        self._message(
            NOM_CHATBOT,
            "Bonjour ! Je suis l'assistant IA de AJI TSTAJI. "
            "Je peux t'aider à préparer ton CV, chercher un stage "
            "ou un emploi dans l'informatique, l'IA et la Data, "
            "et préparer tes candidatures."
        )

        # ----- Suggestions -----
        suggestions = ctk.CTkFrame(self, fg_color="transparent")
        suggestions.grid(row=2, column=0, sticky="ew", padx=12, pady=(0, 6))
        suggestions.grid_columnconfigure((0, 1, 2), weight=1)

        for colonne, texte in enumerate([
            "Aide-moi pour mon CV",
            "Préparer un entretien",
            "Chercher un stage",
        ]):
            ctk.CTkButton(
                suggestions, text=texte, height=30,
                fg_color=BLANC, text_color=BLEU, hover_color="#E8ECFF",
                border_width=1, border_color="#D5DBF5",
                font=ctk.CTkFont(size=11),
                command=lambda t=texte: self._utiliser_suggestion(t),
            ).grid(row=0, column=colonne, sticky="ew", padx=3)

        # ----- Zone de saisie -----
        zone_saisie = ctk.CTkFrame(self, fg_color=BLANC, corner_radius=14)
        zone_saisie.grid(row=3, column=0, sticky="ew", padx=12, pady=(0, 4))
        zone_saisie.grid_columnconfigure(0, weight=1)

        self.champ_message = ctk.CTkTextbox(
            zone_saisie, height=70, wrap="word",
            fg_color="#F4F6FB", text_color=TEXTE, corner_radius=10,
        )
        self.champ_message.grid(
            row=0, column=0, columnspan=3, sticky="ew", padx=8, pady=(8, 4))
        self.champ_message.insert("1.0", PLACEHOLDER)
        self.champ_message.bind("<FocusIn>", self._effacer_placeholder)
        self.champ_message.bind("<Control-Return>", self._ctrl_enter)

        ctk.CTkButton(
            zone_saisie, text="Joindre fichier", width=110, height=30,
            fg_color="transparent", text_color=BLEU,
            border_width=1, border_color=BLEU, hover_color="#E8ECFF",
            font=ctk.CTkFont(size=12), command=self.joindre_fichier,
        ).grid(row=1, column=0, sticky="w", padx=8, pady=(2, 8))

        self.label_fichier = ctk.CTkLabel(
            zone_saisie, text="Aucun fichier", text_color=GRIS,
            font=ctk.CTkFont(size=11), anchor="w")
        self.label_fichier.grid(row=1, column=1, sticky="w", padx=4, pady=(2, 8))

        self.bouton_envoyer = ctk.CTkButton(
            zone_saisie, text="Envoyer", width=90, height=30,
            fg_color=BLEU, hover_color=BLEU_HOVER,
            font=ctk.CTkFont(size=12), command=self.envoyer_message,
        )
        self.bouton_envoyer.grid(row=1, column=2, sticky="e", padx=8, pady=(2, 8))

        # ----- Statut -----
        pied = ctk.CTkFrame(self, fg_color="transparent")
        pied.grid(row=4, column=0, sticky="ew", padx=18, pady=(0, 8))
        pied.grid_columnconfigure(0, weight=1)

        self.label_statut = ctk.CTkLabel(
            pied, text="Assistant prêt", text_color=GRIS,
            font=ctk.CTkFont(size=11), anchor="w")
        self.label_statut.grid(row=0, column=0, sticky="w")

        ctk.CTkLabel(
            pied, text="Ctrl + Entrée pour envoyer", text_color=GRIS,
            font=ctk.CTkFont(size=11), anchor="e",
        ).grid(row=0, column=1, sticky="e")

    def _retour(self):
        if self.fermer:
            self.fermer()

    # ---------------------------------------------------------
    # MESSAGES
    # ---------------------------------------------------------

    def _defiler_bas(self):
        try:
            self.zone_messages.update_idletasks()
            self.zone_messages._parent_canvas.yview_moveto(1.0)
        except Exception:
            pass

    def _message(self, auteur, texte):
        est_utilisateur = auteur == "Vous"

        ligne = ctk.CTkFrame(self.zone_messages, fg_color="transparent")
        ligne.pack(fill="x", pady=5, padx=2)

        if est_utilisateur:
            bulle = ctk.CTkFrame(ligne, fg_color=BULLE_USER, corner_radius=14)
            bulle.pack(side="right", padx=(50, 4))
        else:
            ctk.CTkLabel(ligne, image=self.img_avatar_petit, text="").pack(
                side="left", anchor="n", padx=(2, 6))
            bulle = ctk.CTkFrame(ligne, fg_color=BLANC, corner_radius=14)
            bulle.pack(side="left", fill="x", expand=True, padx=(0, 24))

        ctk.CTkLabel(
            bulle, text=auteur, anchor="w",
            text_color=BLEU if not est_utilisateur else TEXTE,
            font=ctk.CTkFont(size=12, weight="bold"),
        ).pack(fill="x", padx=12, pady=(8, 2))

        affiche = False
        if not est_utilisateur:
            # Réponses de l'IA : Markdown mis en forme
            try:
                contenu = afficher_markdown(
                    bulle,
                    texte,
                    couleur_fond=BLANC,
                    couleur_texte=TEXTE,
                    couleur_accent=BLEU,
                    police="Roboto",
                    taille=10,
                    largeur=34,
                )
                contenu.pack(fill="x", padx=12, pady=(0, 10))
                affiche = True
            except Exception:
                affiche = False

        if not affiche:
            ctk.CTkLabel(
                bulle, text=texte, justify="left", anchor="w",
                text_color=TEXTE, wraplength=340,
            ).pack(fill="x", padx=12, pady=(0, 10))

        self.messages.append((auteur, texte))

        self._defiler_bas()
        self.after(150, self._defiler_bas)

    def _utiliser_suggestion(self, texte):
        self.champ_message.delete("1.0", "end")
        self.champ_message.insert("1.0", texte)
        self.champ_message.focus_set()

    def _effacer_placeholder(self, event=None):
        texte = self.champ_message.get("1.0", "end").strip()
        if texte == PLACEHOLDER:
            self.champ_message.delete("1.0", "end")

    def _ctrl_enter(self, event=None):
        self.envoyer_message()
        return "break"

    # ---------------------------------------------------------
    # ENVOI DU MESSAGE
    # ---------------------------------------------------------

    def envoyer_message(self):
        texte = self.champ_message.get("1.0", "end").strip()

        if texte == PLACEHOLDER:
            texte = ""

        if not texte and not self.fichier_joint:
            self.label_statut.configure(
                text="Écris un message avant de l'envoyer.")
            return

        if not self.api_key:
            self._message(
                NOM_CHATBOT,
                "Clé GROQ_API_KEY introuvable dans le fichier .env.")
            return

        if OpenAI is None:
            self._message(
                NOM_CHATBOT,
                "Bibliothèque manquante. Dans le terminal, "
                "exécute : python -m pip install openai")
            return

        if not self.client:
            try:
                self.client = OpenAI(
                    api_key=self.api_key,
                    base_url="https://api.groq.com/openai/v1",
                )
            except Exception as erreur:
                self._message(NOM_CHATBOT, self._message_erreur(erreur))
                return

        if self.fichier_joint:
            try:
                contenu_fichier = self._lire_fichier(self.fichier_joint)
            except Exception as erreur:
                import traceback
                traceback.print_exc()
                self._message(
                    NOM_CHATBOT,
                    f"Impossible de lire le fichier : "
                    f"{type(erreur).__name__}: {erreur}")
                return

            if not texte:
                texte = (
                    "Analyse ce CV en détail. "
                    "Présente les points forts, les points faibles, "
                    "les erreurs, les compétences, et les améliorations "
                    "à apporter pour candidater à un stage en IA ou Data."
                )

            texte += (
                "\n\n--- CONTENU DU FICHIER À ANALYSER ---\n"
                + contenu_fichier
                + "\n--- FIN DU FICHIER ---"
            )

        self._message("Vous", texte)

        self.champ_message.delete("1.0", "end")
        self.fichier_joint = None
        self.label_fichier.configure(text="Aucun fichier")

        self.bouton_envoyer.configure(state="disabled")
        self.label_statut.configure(text="L'assistant prépare une réponse...")

        threading.Thread(
            target=self._demander_ia, args=(texte,), daemon=True,
        ).start()

    # ---------------------------------------------------------
    # CONNEXION À GROQ
    # ---------------------------------------------------------

    def _nettoyer(self, texte):
        texte = str(texte)
        if self.api_key:
            texte = texte.replace(self.api_key, "***")
        return texte

    def _message_erreur(self, erreur):
        brut = self._nettoyer(erreur)
        t = brut.lower()
        detail = "\nDétail : " + brut[:300]

        if "429" in t or "rate_limit" in t or "quota" in t:
            return ("La limite d'utilisation de Groq est atteinte. "
                    "Attends un peu avant de réessayer." + detail)

        if ("401" in t or "403" in t or "authentication" in t
                or "api_key" in t):
            return ("La clé GROQ_API_KEY est refusée. "
                    "Vérifie sa valeur dans le fichier .env." + detail)

        if "model" in t and ("not found" in t or "decommissioned" in t):
            return ("Le modèle Groq configuré n'est pas disponible. "
                    "Vérifie la variable GROQ_MODEL dans le fichier .env."
                    + detail)

        if any(x in t for x in (
                "connect", "timeout", "timed out", "network", "unreachable")):
            return ("Impossible de joindre Groq. "
                    "Vérifie ta connexion Internet." + detail)

        return "La connexion à Groq a échoué." + detail

    def _demander_ia(self, message):
        try:
            instructions = (
                "Tu es l'assistant IA de AJI TSTAJI, une application "
                "qui aide les étudiants à rechercher des stages et des "
                "emplois, surtout en informatique, intelligence artificielle "
                "et Data à Marrakech. "
                "Réponds en français simple ou en darija marocaine selon "
                "la langue de l'utilisateur. "
                "Donne des conseils concrets et faciles à suivre. "
                "N'invente jamais des offres d'emploi, des entreprises "
                "ou des coordonnées. Si tu ne connais pas une information, "
                "précise-le. Le nom de l'application est AJI TSTAJI."
            )

            historique = [{"role": "system", "content": instructions}]

            for auteur, contenu in self.messages[-10:]:
                if auteur == "Vous":
                    historique.append({"role": "user", "content": contenu})
                elif auteur == NOM_CHATBOT:
                    historique.append({"role": "assistant", "content": contenu})

            if not historique or historique[-1].get("content") != message:
                historique.append({"role": "user", "content": message})

            resultat = self.client.chat.completions.create(
                model=self.modele,
                messages=historique,
                max_completion_tokens=700,
            )

            reponse = resultat.choices[0].message.content
            if not reponse:
                reponse = "Je n'ai pas reçu de réponse. Réessaie."

        except Exception as erreur:
            reponse = self._message_erreur(erreur)

        try:
            self.after(0, self._afficher_reponse, reponse)
        except Exception:
            pass

    def _afficher_reponse(self, reponse):
        self._message(NOM_CHATBOT, reponse)
        self.label_statut.configure(text="Assistant prêt")
        self.bouton_envoyer.configure(state="normal")

    def _lire_fichier(self, chemin):
        """Lit le texte d'un fichier PDF, Word ou TXT."""
        extension = Path(chemin).suffix.lower()

        if extension == ".pdf":
            lecteur = PdfReader(chemin)
            texte = "\n".join(
                page.extract_text() or "" for page in lecteur.pages)

        elif extension == ".docx":
            document = Document(chemin)
            paragraphes = [
                p.text for p in document.paragraphs if p.text.strip()]
            for tableau in document.tables:
                for ligne in tableau.rows:
                    paragraphes.append(
                        " | ".join(cellule.text for cellule in ligne.cells))
            texte = "\n".join(paragraphes)

        elif extension == ".txt":
            with open(chemin, "r", encoding="utf-8-sig") as fichier:
                texte = fichier.read()

        else:
            raise ValueError(
                "Format non pris en charge. Utilise un fichier PDF, DOCX ou TXT.")

        if not texte.strip():
            raise ValueError(
                "Aucun texte n'a été trouvé. "
                "Si ton PDF est scanné comme une image, il faudra utiliser l'OCR.")

        return texte[:18000]

    # ---------------------------------------------------------
    # FICHIER JOINT
    # ---------------------------------------------------------

    def joindre_fichier(self):
        chemin = filedialog.askopenfilename(
            title="Choisir un CV ou un document",
            filetypes=[
                ("Fichiers pris en charge", "*.pdf *.docx *.txt"),
                ("PDF", "*.pdf"),
                ("Word", "*.docx"),
                ("Texte", "*.txt"),
            ],
        )

        if chemin:
            self.fichier_joint = chemin
            nom = os.path.basename(chemin)
            if len(nom) > 22:
                nom = nom[:19] + "..."
            self.label_fichier.configure(text=nom)
            self.label_statut.configure(text="Fichier prêt à être analysé.")

    # ---------------------------------------------------------
    # NOUVELLE CONVERSATION
    # ---------------------------------------------------------

    def nouvelle_conversation(self):
        for widget in self.zone_messages.winfo_children():
            widget.destroy()

        self.messages = []
        self.fichier_joint = None
        self.label_fichier.configure(text="Aucun fichier")
        self.champ_message.delete("1.0", "end")

        self._message(
            NOM_CHATBOT,
            "Nouvelle conversation. Comment puis-je t'aider ?")

        self.label_statut.configure(text="Nouvelle conversation")